# -*- coding: utf-8 -*-
"""Независимая перепроверка расчёта v8: всё заново из исходных данных (не импортирует calc_v6/calc_v8),
другими методами (бисекция вместо перебора, своя матрица по табл. 10, своё моделирование ω(t)),
сверка каждого числа с текстом/таблицами docx и с данными рисунков (npz)."""
import sys, re
import numpy as np
from math import pi, log, sin, cos, tan, atan, radians as rad, degrees as deg, sqrt
from docx import Document

DOCX = sys.argv[1]
out = []
def say(*a):
    out.append(" ".join(str(x) for x in a))

# ------------------------------------------------------------------ исходные данные (табл. 2)
sig_t, kk, Dk, Dobzh, s, n = 415.0, 1.05, 32.10, 28.6, 0.22, 2.0
Pdop, Sz, Sob, S, Lmax, emin, phimax = 120.0, 10.0, 6.0, 16.0, 500.0, 10.0, 60.0
vmax, E = 0.7, 2.1e5
psi_acc, e_acc, L_acc = 36.75, 12.53, 500.0
xM, yM = 50.0, 500.0

def f(x, d=2):
    return f"{x:.{d}f}".replace(".", ",").replace("-", "−")

def bisect(g, a, b, tol=1e-13, it=300):
    ga, gb = g(a), g(b)
    assert ga * gb <= 0, ("нет смены знака", a, b, ga, gb)
    for _ in range(it):
        m = 0.5 * (a + b); gm = g(m)
        if ga * gm <= 0:
            b, gb = m, gm
        else:
            a, ga = m, gm
        if b - a < tol:
            break
    return 0.5 * (a + b)

# ------------------------------------------------------------------ текст docx
doc = Document(DOCX)
norm = lambda t: re.sub(r"\s+", " ", t.replace(" ", " ")).strip()
paras = [norm(p.text) for p in doc.paragraphs]
cells = []
for t in doc.tables:
    for r in t.rows:
        for c in r.cells:
            cells.append(norm(c.text))
FULL = " ¦ ".join(paras + cells)
miss = []
def has(sv, label=""):
    sv = norm(sv)
    ok = sv in FULL
    if not ok:
        miss.append((label, sv))
    return ok

def table(*keys):
    for t in doc.tables:
        head = " ".join(norm(c.text) for c in t.rows[0].cells)
        if all(k in head for k in keys):
            return [[norm(c.text) for c in r.cells] for r in t.rows]
    raise KeyError(keys)

tab_bad = []
def cmp_row(tname, got, exp):
    for j, (g, e_) in enumerate(zip(got, exp)):
        if e_ is None:
            continue
        if g != e_:
            tab_bad.append(f"{tname}: строка «{got[0]}», столбец {j}: в документе «{g}», пересчёт «{e_}»")

# ================================================================== 3. Усилие
sig_rab = kk * sig_t
A_torc = pi / 4 * (26.55 ** 2 - 17.25 ** 2)
F_torc = sig_rab * A_torc
Aring = pi * (Dobzh - s) * s
lnr = log(Dk / Dobzh)
F_uk = n * sig_rab * Aring * lnr
say(f"3: σраб={sig_rab}, Aторц={A_torc:.2f}, F={F_torc/1000:.2f} кН ({F_torc/9.80665/1000:.1f} т, в {F_torc/F_uk:.1f} раз), A={Aring:.4f}, ln={lnr:.5f}, Fук={F_uk:.2f}")
has(f"σраб = k·σт = 1,05·415 = {f(sig_rab, 2)} МПа", "3.1")
has(f"Aторц = π/4·(26,552 − 17,252) = {f(A_torc, 1)} мм2, F = 435,75·319,9 ≈ {f(F_torc / 1000, 0)} кН", "3.2")
has(f"A = π·(28,6 − 0,22)·0,22 = {f(Aring, 2)} мм2, ln(32,10/28,6) = {f(lnr, 4)}", "3.3 A")
has(f"Fук = 2,0·435,75·19,61·0,1154 = {f(F_uk, 0)} Н ≈ 2,0 кН", "3.3 F")
t3 = table("Dобж, мм", "n = 1,5")
exp3 = {}
for D in (28.6, 28.7, 28.8):
    AD = pi * (D - s) * s; lD = log(Dk / D)
    exp3[f(D, 1)] = [f(D, 1), f(AD, 2), f(lD, 4)] + [f(nn * sig_rab * AD * lD, 0) for nn in (1.5, 1.8, 2.0)]
for r in t3[1:]:
    cmp_row("Табл. 3", r, exp3[r[0]])
Fmin3 = 1.5 * sig_rab * pi * (28.8 - s) * s * log(Dk / 28.8)
has(f"{f(Fmin3 / 1000, 2)}–{f(F_uk / 1000, 2)} кН", "табл. 4 диапазон")
say(f"табл. 4: 450 lb = {450 * 4.4482216:.0f} Н; Grifo 120·18,4 = {120 * 18.4:.0f} Н; π·dk·s·σраб = {pi * 26.75 * 0.22 * sig_rab:.0f} Н; 1974/800 = {F_uk / 800:.2f}")
# ================================================================== 4, 5
say("4.1: W = 3·3 − 2·4 =", 3 * 3 - 2 * 4, "; k = 4 − 3 =", 4 - 3)
say(f"табл. 5: 44·0,8·0,5·1·0,5 = {44 * 0.8 * 0.5 * 1.0 * 0.5:.1f}; 75·0,2 = {75 * 0.2:.0f}; 4 кгс = {4 * 9.80665:.1f} Н; 14 кгс = {14 * 9.80665:.1f} Н")
s_ruki = F_uk * Sob / Pdop
has(f"sруки ≥ Fук·Sоб / P = 1974·6 / 120 = {f(s_ruki, 0)} мм", "5.3")
say(f"5.3: путь руки при P = 9…15 Н: {F_uk * Sob / 15 / 1000:.2f}…{F_uk * Sob / 9 / 1000:.2f} м")

# ================================================================== 6. Оптимизация окна (свой алгоритм)
def e_of(pk, psi):
    return S / (tan(rad(pk + psi)) - tan(rad(pk)))
def phob_of(pk, e):
    return deg(atan(tan(rad(pk)) + Sob / e))
F_opt = 1974.0                                     # п. 6.2: Fук = 1974 Н
def L_of(pk, psi):
    e = e_of(pk, psi)
    return F_opt * e / (Pdop * cos(rad(phob_of(pk, e))) ** 3)

def best(psi):
    """Наибольший допустимый φк (A = 8e убывает с ростом φк при φк > −ψ/2): φн ≤ 60, e ≥ 10, L ≤ 500."""
    hi = phimax - psi
    if e_of(hi, psi) < emin:                       # ограничение e ≥ 10
        hi = bisect(lambda pk: e_of(pk, psi) - emin, 0.0, hi)
    if L_of(hi, psi) <= Lmax:
        return hi
    grid = np.linspace(0.0, hi, 20001)             # L не обязательно монотонна — ищем последнюю допустимую точку
    Lg = np.array([L_of(pk, psi) for pk in grid])
    okg = np.where(Lg <= Lmax)[0]
    if okg.size == 0:
        return None
    i = okg[-1]
    return bisect(lambda pk: L_of(pk, psi) - Lmax, grid[i], grid[i + 1])

def Amin(psi):
    pk = best(psi)
    return None if pk is None else 8.0 * e_of(pk, psi)

def minL(psi):
    hi = phimax - psi
    if e_of(hi, psi) < emin:
        hi = bisect(lambda pk: e_of(pk, psi) - emin, 0.0, hi)
    g = np.linspace(0.0, hi, 4001)
    return min(L_of(pk, psi) for pk in g)

psi_star = bisect(lambda p: minL(p) - Lmax, 25.0, 33.0, tol=1e-7)
A_star = 8.0 * emin
psi_kink = bisect(lambda p: L_of(phimax - p, p) - Lmax, 33.0, 45.0)
psi_emin = bisect(lambda p: e_of(phimax - p, p) - emin, 45.0, 59.0)
psi_c = bisect(lambda p: Amin(p) / A_star - p / 29.33, 33.0, psi_kink - 1e-6, tol=1e-9)   # ψ* как в тексте
A_c = Amin(psi_c); e_c = A_c / 8
pk_k = phimax - psi_kink
dA = -S * S / 2 * (1 / cos(rad(pk_k)) ** 2) / (tan(rad(phimax)) - tan(rad(pk_k))) ** 2
elast = -dA * rad(psi_kink) / (8 * e_of(pk_k, psi_kink))
psi_e12 = phimax - deg(atan(tan(rad(phimax)) - S / 12.0))
say(f"6: ψ* = {psi_star:.4f}°, излом ψ = {psi_kink:.3f}°, e = 10 с ψ = {psi_emin:.3f}°, пересечение ψc = {psi_c:.4f}°, Ac = {A_c:.3f}, ec = {e_c:.4f}, δ = {100 * (A_c / A_star - 1):.2f} %")
say(f"   эластичность после излома {elast:.4f}; e = 12 → ψ ≥ {psi_e12:.3f}°; A(30) = {Amin(30):.2f} (δA {100 * (Amin(30) / 80 - 1):.1f} %)")
has(f"ψ ≈ {f(psi_emin, 1)}°, e = 10 мм, A = 80 мм2, но угловой ход при этом на {f(100 * (psi_emin / psi_star - 1), 0)} % больше", "6.6 строгий минимум")
has(f"ψ* = {f(psi_star, 2)}° (наименьший ψ при L ≤ 500 мм)", "6.6 ψ*")
has(f"ψ = {f(psi_c, 2)}°, A = {f(A_c, 1)} мм2, e = {f(e_c, 2)} мм, каждая уступка {f(100 * (A_c / A_star - 1), 1)} %", "6.6 пересечение")
has(f"(при ψ = 30° уступка по площади {f(100 * (Amin(30) / 80 - 1), 0)} %)", "6.6 ψ=30")
has(f"после излома при ψ = {f(psi_kink, 1)}° увеличение ψ на 1 % уменьшает A лишь на {f(elast, 2)} %", "6.6 излом")
has(f"окно с φн ≤ 60° требует ψ ≥ {f(psi_e12, 1)}°", "6.6 e=12")
has(f"при e = 14 мм (Ra20) площадь возрастает до {f(8 * 14, 0)} мм2 (уступка {f(100 * (8 * 14 / 80 - 1), 0)} %)", "6.6 e=14")
# e = 14 при ψ = 36,75: допустимо ли?
pk14 = bisect(lambda pk: e_of(pk, psi_acc) - 14.0, 0.0, phimax - psi_acc)
say(f"   e = 14 при ψ = 36,75: окно {pk14 + psi_acc:.2f} → {pk14:.2f}, L0 = {L_of(pk14, psi_acc):.1f}")

t6 = table("ψ, °", "φоб, °")
for r in t6[1:]:
    psi = float(r[0].replace(",", "."))
    pk = best(psi); e = e_of(pk, psi); A = 8 * e
    if psi == 36.75:
        pk = bisect(lambda x: e_of(x, psi) - 12.53, 0.0, phimax - psi); e = 12.53; A = 8 * e
    exp = [None, f(pk, 2), f(pk + psi, 2), f(e, 2), f(A, 1), f(phob_of(pk, e), 2), f(L_of(pk, psi), 0),
           f(100 * (A / A_star - 1), 1), f(100 * (psi / psi_star - 1), 1)]
    cmp_row("Табл. 6", r, exp)

# ================================================================== 7. Принятое окно
pk = bisect(lambda x: e_of(x, psi_acc) - e_acc, 0.0, phimax - psi_acc)
phi_k, phi_n = pk, pk + psi_acc
phi_ob = phob_of(phi_k, e_acc)
eps_y = sig_t / E; dr_y = eps_y * Dk / 2; tga = ((Dk - Dobzh) / 2) / Sob; d_y = dr_y / tga
phi_y = deg(atan(tan(rad(phi_k)) + (Sob - d_y) / e_acc))
L0 = F_uk * e_acc / (Pdop * cos(rad(phi_ob)) ** 3)
A_acc = 0.5 * e_acc * S
say(f"7: φн = {phi_n:.4f}, φк = {phi_k:.4f}, φоб = {phi_ob:.4f}, φу = {phi_y:.4f}, L0 = {L0:.2f}, A = {A_acc:.3f}, δу = {d_y:.5f}")
has(f"φн = {f(phi_n, 2)}°, φк = {f(phi_k, 2)}°, ψ = 36,75°, φоб = arctg(tg 22,72° + 6/12,53) = {f(phi_ob, 2)}°", "7.1 окно")
has(f"L0 = 1974·12,53 / (120·cos341,91°) ≈ {f(L0, 0)} мм", "7.1 L0")
has(f"A = ½·12,53·16 = {f(A_acc, 1)} мм2", "7.1 A")

# ================================================================== 11. Закон силы, система 9 уравнений, L_треб
A_pr = n * Aring * lnr
F_t = sig_t * A_pr; F_max = kk * sig_t * A_pr
C = F_t / d_y; k_h = (F_max - F_t) / (Sob - d_y)
def Fphi(phi):
    if phi > phi_ob:
        return 0.0
    dl = e_acc * (tan(rad(phi_ob)) - tan(rad(phi)))
    return C * dl if dl <= d_y else F_t + k_h * (dl - d_y)
say(f"11.2: Aпр = {A_pr:.4f}, Fт = {F_t:.2f}, Fук = {F_max:.2f}, C = {C:.1f} (по округлённым 1880/0,109 = {1880 / 0.109:.0f}), k = {k_h:.3f}, ΔF = {F_max - F_t:.2f}")
has(f"Aпр = n·A·ln(Dk/Dобж) = 2,0·19,61·0,1154 = {f(A_pr, 3)} мм2", "11.2 Aпр")
has(f"Fт = σт·Aпр = 415·4,529 = {f(F_t, 0)} Н", "11.2 Fт")
has(f"Fук = 1,05·σт·Aпр = 435,75·4,529 = {f(F_max, 0)} Н", "11.2 Fук")
has(f"k = (Fук − Fт)/(Sоб − δу) = {f(F_max - F_t, 0)}/{f(Sob - d_y, 3)} ≈ {f(k_h, 1)} Н/мм", "11.2 k")
has(f"F(φ) = 0, φ > φоб = {f(phi_ob, 2)}°;", "11.2 закон")
has(f"φк ≤ φ < φу = {f(phi_y, 2)}°;", "11.2 закон φу")
has(f"Упругий участок соответствует повороту кулисы всего на {f(phi_ob - phi_y, 2)}°", "11.2 0,28")

def system(phi, L=L_acc, xm=xM, ym=yM):
    """Матрица ровно по табл. 10 документа (неизвестные R10x, R10y, R12N, M12, R23x, R23y, R30N, M30, P)."""
    F = Fphi(phi); p = rad(phi)
    XD, YD, XK = xm + e_acc, ym - e_acc * tan(p), xm + L * cos(p)
    a = XD * sin(p + pi / 2) + YD * cos(p + pi / 2)
    A = np.array([
        [1, 0, cos(p + pi / 2), 0, 0, 0, 0, 0, 0],
        [0, 1, sin(p + pi / 2), 0, 0, 0, 0, 0, -1],
        [ym, xm, a, 1, 0, 0, 0, 0, -XK],
        [0, 0, cos(p - pi / 2), 0, 1, 0, 0, 0, 0],
        [0, 0, sin(p - pi / 2), 0, 0, 1, 0, 0, 0],
        [0, 0, -a, -1, YD, XD, 0, 0, 0],
        [0, 0, 0, 0, -1, 0, 1, 0, 0],
        [0, 0, 0, 0, 0, -1, 0, 0, 0],
        [0, 0, 0, 0, -YD, -XD, YD, 1, 0]], dtype=float)
    B = np.array([0, 0, 0, 0, 0, 0, 0, -F, -F * XD], dtype=float)
    return A, B, (XD, YD, XK, a), F

phis = np.linspace(phi_k, phi_n, 40001)
dets, maxdev, maxM = [], 0.0, 0.0
for ph in phis:
    A_, B_, _, F = system(ph)
    X = np.linalg.solve(A_, B_)
    dets.append(abs(np.linalg.det(A_)))
    p = rad(ph)
    an = np.array([F * tan(p), F * e_acc / (L_acc * cos(p) ** 3) - F, F / cos(p), 0, -F * tan(p), F, -F * tan(p), 0,
                   F * e_acc / (L_acc * cos(p) ** 3)])
    maxdev = max(maxdev, float(np.max(np.abs(X - an))))
    X2 = np.linalg.solve(*system(ph, xm=-37.0, ym=811.0)[:2])        # другая точка M
    maxM = max(maxM, float(np.max(np.abs(X - X2))))
say(f"11.5: |det A| = {min(dets):.1f}…{max(dets):.1f}; отклонение от аналитики {maxdev:.2e} Н; влияние точки M {maxM:.2e} Н")
# пример при φу
A_, B_, (XD, YD, XK, a), F = system(phi_y)
X = np.linalg.solve(A_, B_)
say(f"   при φу: XD={XD:.3f} YD={YD:.3f} XK={XK:.3f} a={a:.3f}; X = " + ", ".join(f"{v:.2f}" for v in X))
has(f"XD = {f(XD, 2)} мм, YD = {f(YD, 2)} мм, XK = {f(XK, 2)} мм, a(φ) = {f(a, 2)} мм", "11.5 пример")
has(f"R10x = {f(X[0], 1)} Н, R10y = {f(X[1], 1)} Н, R12N = {f(X[2], 1)} Н, M12 = 0, R23x = {f(X[4], 1)} Н, "
    f"R23y = {f(X[5], 1)} Н, R30N = {f(X[6], 1)} Н, M30 = 0, P = {f(X[8], 1)} Н", "11.5 решение")
has(f"(|det A| = {f(min(dets), 0)}…{f(max(dets), 0)})", "11.5 det")
# невязки уравнений при φу с числами из документа
Xdoc = np.array([1670.7, -1766.7, 2514.8, 0, -1670.7, 1879.6, -1670.7, 0, 112.8])
res = A_ @ Xdoc - B_
say("   невязки 9 уравнений на числах документа (Н, Н·мм):", ", ".join(f"{v:.1f}" for v in res))
# L_треб(φ)
Lreq = np.array([Fphi(ph) * e_acc / (Pdop * cos(rad(ph)) ** 3) for ph in phis])
i = int(np.argmax(Lreq)); Lreq_max, ph_Lmax = Lreq[i], phis[i]
Pmax = Pdop * Lreq_max / L_acc
say(f"11.7: Lтреб max = {Lreq_max:.3f} при φ = {ph_Lmax:.4f}; P max = {Pmax:.3f}; запас {(1 - Pmax / Pdop) * 100:.1f} %")
has(f"Наибольшее значение Lтреб max = {f(Lreq_max, 1)} мм достигается в начале пластической деформации (φ = {f(ph_Lmax, 2)}°)", "11.7 Lmax")
has(f"P = 120·470,1/500 = {f(Pmax, 1)} Н ≤ 120 Н", "11.7 P")
Mb = Pmax * (L_acc * cos(rad(phi_y)) - e_acc) / 1000
say(f"   изгибающий момент у камня {Mb:.4f} Н·м")
has(f"M = P·(L·cos φ − e) ≈ {f(Mb, 1)} Н·м", "11.7 M")
# табл. 11
t11 = table("Положение", "|R10|")
pos = {"начало хода": phi_n, "касание с крышкой": phi_ob, "начало пластической деформации": phi_y, "промежуточное": 30.0, "конец хода": phi_k}
for r in t11[1:]:
    ph = pos[r[0]]
    A_, B_, _, F = system(ph); X = np.linalg.solve(A_, B_)
    X = np.where(np.abs(X) < 1e-9, 0.0, X)
    exp = [None, f(ph, 2) + "°", f(F, 0), f(X[8], 1) if F > 0 else "0", f(X[0], 0), f(X[1], 0), f(np.hypot(X[0], X[1]), 0),
           f(X[2], 0), f(X[4], 0), f(X[5], 0), f(X[6], 0)]
    cmp_row("Табл. 11", r, exp)
Xend = np.linalg.solve(*system(phi_k)[:2])
has(f"|R10y| = {f(abs(Xend[1]), 0)} Н — в конце хода (φ = {f(phi_k, 2)}°)", "11.7 R10y")
has(f"P = {f(Pmax, 1)} Н, реакция в паре кулиса–камень R12N = {f(np.linalg.solve(*system(phi_y)[:2])[2], 0)} Н", "11.7 итог")
# 11.1 оценка рукояти
m = 7850 * 0.020 * 0.008 * 0.5; J = m * 0.5 ** 2 / 3
say(f"11.1: m = {m:.3f} кг, J = {J:.4f}, вес {m * 9.81:.2f} Н → {m * 9.81 / 2:.2f} Н руки")

# ================================================================== 8. Аналоги
def analogs(ph, L=L_acc):
    p = rad(ph); c, s_ = cos(p), sin(p)
    return dict(lOD=e_acc / c, VqDy=-e_acc / c ** 2, aqDy=2 * e_acc * s_ / c ** 3, Vqr=e_acc * s_ / c ** 2,
                aqr=e_acc * (1 + s_ ** 2) / c ** 3, FP=L * c ** 3 / e_acc, VqKx=L * s_, VqKy=-L * c)
t8 = table("Положение", "lOD")
pos8 = {"начало хода": phi_n, "начало обжима": phi_ob, "конец хода": phi_k}
for r in t8[1:]:
    ph = pos8[r[0]]; q = analogs(ph)
    cmp_row("Табл. 8", r, [None, f(ph, 2) + "°", f(q["lOD"], 2), f(q["VqDy"], 2), f(q["aqDy"], 2), f(q["Vqr"], 2),
                           f(q["aqr"], 2), f(2 * q["Vqr"], 2), f(q["FP"], 2)])
q = analogs(phi_ob)
has(f"VqKx = L·sin φ = {f(q['VqKx'], 1)} мм/рад, VqKy = −L·cos φ = {f(q['VqKy'], 1)} мм/рад", "8 K")
FPy = analogs(phi_y)["FP"]
has(f"(φу = {f(phi_y, 2)}°) оно равно {f(FPy, 2)}, и сила руки 1880/16,66 = {f(F_t / FPy, 1)} Н", "8 F/P")

# ================================================================== 9. Закон движения (своё моделирование)
psi_xx, psi_u, psi_pl = rad(phi_n - phi_ob), rad(phi_ob - phi_y), rad(phi_y - phi_k)
w = vmax / (L_acc / 1000)
t_r, t_u, t_t = 2 * psi_xx / w, psi_u / w, 2 * psi_pl / w
eps_r, eps_t = w / t_r, -w / t_t
T = t_r + t_u + t_t
say(f"9: ψхх = {deg(psi_xx):.3f}° = {psi_xx:.5f}, ψу = {deg(psi_u):.4f}° = {psi_u:.6f}, ψпл = {deg(psi_pl):.3f}° = {psi_pl:.5f}")
say(f"   ω = {w:.4f}, tр = {t_r:.5f}, tу = {t_u * 1000:.4f} мс, tт = {t_t:.5f}, εр = {eps_r:.4f}, εт = {eps_t:.4f}, T = {T:.5f}")
has(f"εт = σт/E = 415 / 2,1·105 = {f(eps_y, 5)}", "9.2 εт")
has(f"Δrу = εт·Dk/2 = 0,00198·32,10/2 = {f(dr_y, 3)} мм", "9.2 Δr")
has(f"tg α = 1,75/6 = {f(tga, 3)} (α = {f(deg(atan(tga)), 1)}°), δу = Δrу/tg α = 0,032/0,292 = {f(d_y, 3)} мм", "9.2 δу")
has(f"φоб = arctg(tg 22,72° + 6/12,53) = {f(phi_ob, 2)}°, φу = arctg(tg 22,72° + 5,891/12,53) = {f(phi_y, 2)}°", "9.2 углы")
has(f"ψхх = 59,47° − 41,91° = {f(deg(psi_xx), 2)}° = {f(psi_xx, 4)} рад", "9.2 ψхх")
has(f"ψу = 41,91° − 41,63° = {f(deg(psi_u), 2)}° = {f(psi_u, 5)} рад, ψпл = 41,63° − 22,72° = {f(deg(psi_pl), 2)}° = {f(psi_pl, 4)} рад", "9.2 ψу ψпл")
has(f"ωmax = vmax/L = 0,7/0,5 = {f(w, 2)} рад/с", "9.3 ω")
has(f"tр = 2ψхх/ωmax = 2·0,3065/1,40 = {f(t_r, 3)} с", "9.3 tр")
has(f"tу = ψу/ωmax = 0,00483/1,40 = {f(t_u, 4)} с", "9.3 tу")
has(f"tт = 2ψпл/ωmax = 2·0,3301/1,40 = {f(t_t, 3)} с", "9.3 tт")
has(f"εр = 1,40/0,438 = {f(eps_r, 2)} рад/с2, εт = −1,40/0,472 = {f(eps_t, 2)} рад/с2", "9.3 ε")
has(f"tрх = 0,438 + 0,003 + 0,472 = {f(T, 3)} с", "9.3 tрх")
has(f"= {f(psi_xx + psi_u + psi_pl, 3)} рад = {f(deg(psi_xx + psi_u + psi_pl), 2)}°", "9.3 проверка")
sK = L_acc * rad(psi_acc)
has(f"sK/tрх = {f(sK, 0)}/0,913 = {f(sK / T, 0)} мм/с", "9.3 vср")
has(f"Установившийся участок короткий ({f(t_u * 1000, 1)} мс)", "9.3 мс")
say(f"9.4: tрх + tобр = {2 * T:.3f} с, остаётся {10 - 2 * T:.2f} с")

# ------------------------------------------------------------------ моделирование по времени
tt = np.linspace(0, T, 400001)
def state(t):
    if t <= t_r:
        return eps_r * t, eps_r, 0.5 * eps_r * t * t
    if t <= t_r + t_u:
        return w, 0.0, psi_xx + w * (t - t_r)
    tau = t - t_r - t_u
    return w + eps_t * tau, eps_t, psi_xx + psi_u + w * tau + 0.5 * eps_t * tau * tau
om, ep, qq = np.array([state(t) for t in tt]).T
ph_t = phi_n - np.degrees(qq)
assert abs(ph_t[-1] - phi_k) < 1e-6 and abs(om[-1]) < 1e-9
c_ = np.cos(np.radians(ph_t)); s_ = np.sin(np.radians(ph_t))
VDy = om * (-e_acc / c_ ** 2)
aDy = om ** 2 * (2 * e_acc * s_ / c_ ** 3) + ep * (-e_acc / c_ ** 2)
Vr = om * e_acc * s_ / c_ ** 2
iv, ir = int(np.argmin(VDy)), int(np.argmax(Vr))
say(f"10: VDy min = {VDy[iv]:.3f} при φ = {ph_t[iv]:.3f}, t = {tt[iv]:.4f}; aDy(0) = {aDy[0]:.3f}; Vr max = {Vr[ir]:.3f} при φ = {ph_t[ir]:.3f}")
has(f"VDy = {f(VDy[iv], 1)} мм/с достигается на холостом ходе при φ = {f(ph_t[iv], 1)}° (t = {f(tt[iv], 3)} с)", "10 Vmax")
Vc = w * analogs(phi_ob)["VqDy"]; Vrc = w * analogs(phi_ob)["Vqr"]
has(f"головка подходит к крышке с VDy = {f(Vc, 1)} мм/с", "10 Vконт")
has(f"aDy = {f(aDy[0], 1)} мм/с2 — в начале хода", "10 a0")
has(f"со скоростью до {f(Vr[ir], 1)} мм/с (на холостом ходе, при φ ≈ {f(ph_t[ir], 1)}°)", "10 Vr")
has(f"Vr = ω·Vq,r = 1,40·15,11 = {f(Vrc, 2)} мм/с; ускорение Кориолиса в момент касания aк = 2·1,40·21,16 = {f(2 * w * Vrc, 1)} мм/с2", "10 Кориолис")
# табл. 9
t9 = table("Момент", "ω, рад/с")
def row9(t, ph, om_, ep_):
    q = analogs(ph)
    V = om_ * q["VqDy"]; a_ = om_ ** 2 * q["aqDy"] + ep_ * q["VqDy"]
    sg = lambda x, d: ("+" if x > 0 else "") + f(x, d)
    return [None, "0" if t == 0 else f(t, 3), f(ph, 2) + "°", "0" if om_ == 0 else f(om_, 2), "0" if ep_ == 0 else sg(ep_, 2),
            "0" if om_ == 0 else f(V, 1), sg(a_, 1), "0" if om_ == 0 else f(om_ * L_acc / 1000, 2),
            "0" if ep_ == 0 else sg(ep_ * L_acc / 1000, 2), "0" if om_ == 0 else f(om_ ** 2 * L_acc / 1000, 2)]
exp9 = [row9(0, phi_n, 0, eps_r), row9(t_r, phi_ob, w, eps_r), row9(t_r, phi_ob, w, 0), row9(t_r + t_u, phi_y, w, 0),
        row9(t_r + t_u, phi_y, w, eps_t), row9(T, phi_k, 0, eps_t)]
for r, e_ in zip(t9[1:], exp9):
    cmp_row("Табл. 9", r, e_)

# ================================================================== 13. Нелинейность
ph_g = np.linspace(phi_k, phi_n, 200001)
yv = e_acc * np.tan(np.radians(ph_g))
chord = yv[0] + (yv[-1] - yv[0]) * (ph_g - ph_g[0]) / (ph_g[-1] - ph_g[0])
dev = float(np.max(np.abs(chord - yv)))
ph_o = np.linspace(0, 30, 200001); e_o = S / tan(rad(30)); yo = e_o * np.tan(np.radians(ph_o))
dev_o = float(np.max(np.abs(yo[0] + (yo[-1] - yo[0]) * ph_o / 30 - yo)))
has(f"от хорды — {f(dev, 2)} мм = {f(100 * dev / S, 1)} % хода", "13 нелинейность")
has(f"для окна 30° → 0° было {f(100 * dev_o / S, 1)} %", "13 старое окно")

# ================================================================== 12. Сводка
has(f"P(φ) ≤ {f(Pmax, 1)} Н при L = 500 мм; допустимая — 120 Н", "12 P")
has(f"{f(phi_n, 2)}° → {f(phi_k, 2)}°", "12 окно")
has(f"500 мм (Lтреб max = {f(Lreq_max, 0)} мм)", "12 L")
has(f"{f(sK, 0)} / {f(L_acc * sin(rad(phi_n)), 0)} мм", "12 дуга/подъём")
has(f"{f(d_y, 3)} / {f(Sob - d_y, 3)} мм", "12 упр/пласт")
has(f"{f(T, 3)} с: {f(t_r, 3)} / {f(t_u, 3)} / {f(t_t, 3)} с", "12 время")
has(f"{f(w, 2)} рад/с; {f(eps_r, 2)}; {f(eps_t, 2)} рад/с2", "12 ω ε")
has(f"{f(-VDy[iv], 1)} мм/с / {f(w * L_acc / 1000, 2)} м/с", "12 V")
Xy = np.linalg.solve(*system(phi_y)[:2])
has(f"{f(Xy[2], 0)} / {f(np.hypot(Xy[0], Xy[1]), 0)} / {f(abs(Xy[6]), 0)} Н (φ = {f(phi_y, 2)}°)", "12 реакции")
has(f"L·sin φн = {f(L_acc * sin(rad(phi_n)), 0)} мм", "11.3 подъём")


# ================================================================== правки v9, табл. 7, ссылки
has("Рабочая версия от 11.10.2026", "v9 дата")
has("C = Fт/δу = 1880/0,1087 ≈ 17 300 Н/мм", "v9 C")
has("F_ук — сила".replace("F_ук", "Fук") + ", с которой головка давит на крышку", "v9 рис. 1")
say(f"C по 1880/0,1087 = {1880 / 0.1087:.0f}")
t7 = table("Размер", "Ra10, мм")
for r in t7[1:]:
    if r[0].startswith("Расстояние e"):
        cmp_row("Табл. 7", r, [None, f(e_acc, 2), "16", "14", f(e_acc, 2), "—", None])
    if r[0].startswith("Длина рукояти L"):
        cmp_row("Табл. 7", r, [None, f"{Lreq_max:.0f} (max Lтреб, п. 11.7)", "500", "500", "500", "Ra10", f"наибольшая сила руки {f(Pmax, 1)} Н ≤ 120 Н"])
heads = set(re.findall(r"^(\d+(?:\.\d+)?)\.\s", "\n".join(paras), flags=re.M))
refs = re.findall(r"(?:пп?\.|разд\.|раздел[еау]?|разделы)\s*((?:\d+(?:\.\d+)?(?:[–,]\s*)?)+)", FULL)
bad_ref = []
for grp in refs:
    for num in re.findall(r"\d+(?:\.\d+)?", grp):
        if num not in heads and not (num.isdigit() and int(num) <= 16):
            bad_ref.append(num)
tabs = [int(x) for x in re.findall(r"[Тт]абл(?:ица|ице|ицы|ицу|\.)\s*(\d+)", FULL)]
figs = [int(x) for x in re.findall(r"[Рр]ис(?:унок|унке|унка|унки|\.)\s*(\d+)", FULL)]
say(f"ссылки: пунктов-заголовков {len(heads)}; ссылок на несуществующие пункты: {sorted(set(bad_ref)) or 'нет'}; "
    f"таблицы {min(tabs)}…{max(tabs)} (всего в документе {len(doc.tables)}), рисунки {min(figs)}…{max(figs)} (всего {len(doc.inline_shapes)})")

# ================================================================== данные рисунков (npz) против независимого расчёта
try:
    fz = np.load("forces_v8.npz"); kz = np.load("kin_v8.npz"); az = np.load("amin_v6.npz")
    Xf, phf = fz["X"], fz["phi"]
    idx = np.linspace(0, len(phf) - 1, 400).astype(int)
    d_forces = max(float(np.max(np.abs(Xf[i] - np.linalg.solve(*system(float(phf[i]))[:2])))) for i in idx)
    d_L = max(abs(float(fz["Lreq"][i]) - Fphi(float(phf[i])) * e_acc / (Pdop * cos(rad(float(phf[i]))) ** 3)) for i in idx)
    tk = kz["t"]; jj = np.linspace(0, len(tk) - 1, 400).astype(int)
    d_om = max(abs(float(kz["om"][j]) - state(float(tk[j]))[0]) for j in jj)
    d_V = max(abs(float(kz["V"][j]) - np.interp(float(tk[j]), tt, VDy)) for j in jj)
    ps, Aa = az["ps"], az["A"]
    sel = [k for k in np.linspace(0, len(ps) - 1, 60).astype(int) if ps[k] >= psi_star + 0.05]
    d_A = max(abs(float(Aa[k]) - Amin(float(ps[k]))) for k in sel)
    say(f"рисунки: реакции (рис. 9) Δ = {d_forces:.2e} Н; Lтреб (рис. 8) Δ = {d_L:.2e} мм; ω(t) (рис. 4) Δ = {d_om:.2e}; "
        f"VDy (рис. 5) Δ = {d_V:.2e} мм/с; Amin(ψ) (рис. 3) Δ = {d_A:.3f} мм² (сетка npz)")
    say(f"   npz: ψ* = {float(az['psi_star']):.4f}, ψc = {float(az['psi_c']):.4f}, Ac = {float(az['A_c']):.3f}, ec = {float(az['e_c']):.4f}; "
        f"φ(npz) н/об/у/к = {float(fz['phi_n']):.3f}/{float(fz['phi_ob']):.3f}/{float(fz['phi_y']):.3f}/{float(fz['phi_k']):.3f}")
except FileNotFoundError as ex:
    say("npz не найдены:", ex)

print("\n".join(out))
print("\n=== Таблицы: расхождений", len(tab_bad))
for b in tab_bad:
    print("  ", b)
print("=== Текст: не найдено строк", len(miss))
for lab, sv in miss:
    print(f"   [{lab}] {sv}")
