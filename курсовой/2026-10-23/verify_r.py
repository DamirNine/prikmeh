# -*- coding: utf-8 -*-
"""Независимая перепроверка раздела 12 с камнем-роликом: заново из исходных данных, без calc_r.py; местный изгиб стенки —
по формулам балки с заделками для сосредоточенной силы; сверка с текстом docx; поиск следов прежнего камня."""
import sys, re
from math import pi, sqrt, tan, atan, cos, sin, radians as rad, degrees as deg, log
from docx import Document

DOCX = sys.argv[1]
bad = []
def chk(name, got, ref, tol):
    ok = abs(got - ref) <= tol
    print(f"{'OK ' if ok else 'ERR'} {name}: {got:.4f} (ожидалось {ref})")
    if not ok:
        bad.append(name)

e, psi, S, Sob, L = 12.53, 36.75, 16.0, 6.0, 500.0
A_pr = 2.0 * pi * (28.6 - 0.22) * 0.22 * log(32.10 / 28.6)
Ft, Fk = 415.0 * A_pr, 1.05 * 415.0 * A_pr
lo, hi = 0.0, 23.25
for _ in range(200):
    m = (lo + hi) / 2
    if 16.0 / (tan(rad(m + psi)) - tan(rad(m))) - e > 0: lo = m
    else: hi = m
phk = (lo + hi) / 2; phn = phk + psi
dy = (415.0 / 2.1e5) * 32.10 / 2 / ((32.10 - 28.6) / 2 / Sob)
phy = deg(atan(tan(rad(phk)) + (Sob - dy) / e))
def at(phi, F):
    c, s = cos(rad(phi)), sin(rad(phi)); P = F * e / (L * c ** 3)
    return dict(F=F, P=P, R12=F / c, R10=sqrt((F * tan(rad(phi))) ** 2 + (P - F) ** 2), N=P * s, Qdk=P * c, MD=P * (L * c - e), lOD=e / c)
A = at(phy, Ft); B = at(phk, Fk)

# контакт ролика: Герц, линия длиной b = 12, ρ = 11; допускаемое по ГОСТ 21354-87 (HRC 45…50 → 47,5; S_H = 1,2)
E = 2.1e5; b = 12.0
sHP = (17 * 47.5 + 200) / 1.2
sH = 0.418 * sqrt(A["R12"] * E / (b * 11.0))
chk("[σH], МПа", sHP, 839.6, 0.05); chk("σH ролика Ø22, МПа", sH, 836.0, 0.6)
print(("OK " if 0.418 * sqrt(A["R12"] * E / (b * 10.0)) > sHP else "ERR"), "Ø20 не проходит:", round(0.418 * sqrt(A["R12"] * E / (b * 10.0)), 1))
Fov = 120 * L * cos(rad(phk)) ** 3 / e
chk("σH при перегрузке", 0.418 * sqrt(Fov / cos(rad(phk)) * E / (b * 11.0)), 1064.0, 0.6)
D_soft = 2 * 0.418 ** 2 * A["R12"] * E / (b * ((2 * 350 + 70) / 1.1) ** 2)
chk("ролик без закалки (350 НВ), мм", D_soft, 31.4, 0.05)

# пруток (сталь 45, [σ] = 142): плечи 7 и 23 мм
chk("палец D Ø11, МПа", 32 * A["R12"] * 7 / (pi * 11 ** 3), 134.7, 0.06)
chk("ось O Ø16, МПа", 32 * A["R10"] * 23 / (pi * 16 ** 3), 139.1, 0.06)
chk("шток Ø12, МПа", 4 * Fk / (pi * 144), 17.4, 0.05)

# рукоять 12 × 12, паз Ø22 от a1 до a2
lmin, lmax = e / cos(rad(phk)), e / cos(rad(phn))
a1, a2 = lmin - 11 - 1, lmax + 11 + 1; ell = a2 - a1
Mh = A["Qdk"] * (L - a2)
chk("рукоять 12×12, МПа", A["N"] / (b * 12) + 6 * Mh / (b * 144), 136.2, 0.05)
def wall(pos):
    c = pos["lOD"] - a1; d = ell - c; F = pos["R12"]
    MA, MB, Mc = -F * c * d ** 2 / ell ** 2, -F * c ** 2 * d / ell ** 2, 2 * F * c ** 2 * d ** 2 / ell ** 3
    return MA, MB, Mc
MA, MB, Mc = wall(A)
chk("M_A, Н·м", MA / 1000, -12.29, 0.006); chk("M_B, Н·м", MB / 1000, -9.37, 0.006); chk("M_c, Н·м", Mc / 1000, 10.63, 0.006)
def sec(H, pos):
    t = (H - 22) / 2; W = b * (H ** 3 - 22 ** 3) / (6 * H)
    MA_, MB_, Mc_ = wall(pos); Ml = max(abs(MA_), abs(MB_), abs(Mc_))
    return pos["N"] / (b * (H - 22)) + pos["MD"] / W + 6 * Ml / (b * t ** 2)
chk("сечение А–А H = 40, φу", sec(40, A), 91.5, 0.06); chk("сечение А–А H = 40, φк", sec(40, B), 79.3, 0.06)
lo, hi = 23.0, 100.0
for _ in range(200):
    m = (lo + hi) / 2
    if max(sec(m, A), sec(m, B)) > 142: lo = m
    else: hi = m
chk("H_треб", lo, 36.2, 0.05)

# текст документа
doc = Document(DOCX)
FULL = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
need = ["836", "840", "1064", "2090", "134,7", "139,1", "136,2", "91,5", "15,6", "75,8", "12,29", "10,63", "9,37",
        "31,4", "Ø22 × 12", "2 — камень (ролик на пальце D)", "0,850", "0,0477", "3758", "1,57", "Таблица 19"]
miss = [s for s in need if s not in FULL]
print("в тексте не найдено:", miss or "—"); bad += miss
traces = [s for s in ("сухар", "брусок", "16 × 16", "18 × 18", "квадратный", "квадратного", "Ст3сп, [σ]", "как просили",
                      "выбрали", "вместо квадрат", "круглый камень", "круглым", "предыдущ", "прошл", "в отличие от") if s in FULL]
print("следы прежнего камня / объяснения:", traces or "—"); bad += traces
print("ИТОГ:", "расхождений нет" if not bad else f"расхождения: {bad}")
