# -*- coding: utf-8 -*-
"""Независимая перепроверка раздела 12 (v11): всё заново из исходных данных, без calc_v11.py;
местный изгиб стенки паза — по аналитическим интегралам формул балки с заделками; сверка чисел с текстом и таблицами docx."""
import sys, re
from math import pi, sqrt, tan, atan, cos, sin, radians as rad, degrees as deg, log
from docx import Document

DOCX = sys.argv[1]
bad = []
def chk(name, got, ref, tol):
    ok = abs(got - ref) <= tol
    print(f"{'OK ' if ok else 'ERR'} {name}: {got:.4f} (ожидалось {ref})")
    if not ok:
        bad.append(name)

# ------------------------------------------------------------------ механизм и закон обжима (разд. 3, 11.2)
e, psi, S, Sob, L = 12.53, 36.75, 16.0, 6.0, 500.0
A_pr = 2.0 * pi * (28.6 - 0.22) * 0.22 * log(32.10 / 28.6)
Ft, Fk = 415.0 * A_pr, 1.05 * 415.0 * A_pr
lo, hi = 0.0, 23.25
for _ in range(200):                                   # φк из условия хода S = e·(tg φн − tg φк)
    m = (lo + hi) / 2
    if 16.0 / (tan(rad(m + psi)) - tan(rad(m))) - e > 0: lo = m
    else: hi = m
phk = (lo + hi) / 2; phn = phk + psi
dy = (415.0 / 2.1e5) * 32.10 / 2 / ((32.10 - 28.6) / 2 / Sob)
phy = deg(atan(tan(rad(phk)) + (Sob - dy) / e))
def at(phi, F):
    c, s = cos(rad(phi)), sin(rad(phi))
    P = F * e / (L * c ** 3)
    return dict(F=F, P=P, R12=F / c, R10=sqrt((F * tan(rad(phi))) ** 2 + (P - F) ** 2), N=P * s, Qdk=P * c, MD=P * (L * c - e),
                lOD=e / c)
A = at(phy, Ft); B = at(phk, Fk)
chk("φу", phy, 41.634, 0.001); chk("φк", phk, 22.720, 0.001)
chk("M_D(φу), Н·м", A["MD"] / 1000, 40.75, 0.005); chk("R12N(φу)", A["R12"], 2514.8, 0.1); chk("|R10|(φу)", A["R10"], 2431.6, 0.1)

# ------------------------------------------------------------------ детали из прутка (сталь 45, [σ] = 142)
sD = 32 * A["R12"] * 5.0 / (pi * 10 ** 3); sO = 32 * A["R10"] * 21.0 / (pi * 16 ** 3); sR = 4 * Fk / (pi * 12 ** 2)
chk("палец D Ø10, МПа", sD, 128.1, 0.05); chk("ось O Ø16, МПа", sO, 127.0, 0.05); chk("шток Ø12, МПа", sR, 17.4, 0.05)
chk("τ пальца D", 16 * A["R12"] / (3 * pi * 100), 42.7, 0.05)
for nm, d, sa in (("палец D: Ø9 не проходит (45)", 9, 142), ("ось O: Ø15 не проходит (45)", 15, 142)):
    s_ = 32 * (A["R12"] * 5.0 if "палец" in nm else A["R10"] * 21.0) / (pi * d ** 3)
    print(("OK " if s_ > sa else "ERR"), nm, round(s_, 1)); bad += [] if s_ > sa else [nm]

# ------------------------------------------------------------------ рукоять и зона паза (лист Ст3сп 8 мм, [σ] = 98)
b, w, lk, g = 8.0, 16.0, 16.0, 1.0
lmin, lmax = e / cos(rad(phk)), e / cos(rad(phn))
a1, a2 = lmin - lk / 2 - g, lmax + lk / 2 + g; ell = a2 - a1
Mh = A["Qdk"] * (L - a2); sh = A["N"] / (b * 18) + 6 * Mh / (b * 18 ** 2)
hreq = (A["N"] + sqrt(A["N"] ** 2 + 24 * 98 * b * Mh)) / (2 * 98 * b)
chk("рукоять 8×18, МПа", sh, 91.5, 0.05); chk("h_треб", hreq, 17.4, 0.05)

def MA_partial(x1, x2, q, Lb):
    """момент в левой заделке от равномерной нагрузки q на [x1, x2]: −q/L²·∫x(L−x)² dx (аналитически)"""
    F = lambda x: Lb ** 2 * x ** 2 / 2 - 2 * Lb * x ** 3 / 3 + x ** 4 / 4
    return -q / Lb ** 2 * (F(x2) - F(x1))
def MB_partial(x1, x2, q, Lb):
    F = lambda x: Lb * x ** 3 / 3 - x ** 4 / 4
    return -q / Lb ** 2 * (F(x2) - F(x1))
def wall(pos, H):
    c = pos["lOD"] - a1; q = pos["R12"] / lk
    MA, MB = MA_partial(c - lk / 2, c + lk / 2, q, ell), MB_partial(c - lk / 2, c + lk / 2, q, ell)
    # момент под нагрузкой (наибольший положительный): M(x) = MA + RA·x − q(x − x1)²/2, RA из равновесия
    x1, x2 = c - lk / 2, c + lk / 2
    RA = (pos["R12"] * (ell - c) + MB - MA) / ell
    xm = x1 + RA / q
    Mmid = MA + RA * xm - q * (xm - x1) ** 2 / 2
    Ml = max(abs(MA), abs(MB), abs(Mmid))
    t = (H - w) / 2; W = b * (H ** 3 - w ** 3) / (6 * H)
    return dict(MA=MA, MB=MB, Mmid=Mmid, s=pos["N"] / (b * (H - w)) + pos["MD"] / W + 6 * Ml / (b * t ** 2))
WA = wall(A, 36.0); WB = wall(B, 36.0)
chk("местный момент в заделке у a1, Н·м", -WA["MA"] / 1000, 8.97, 0.006)
chk("местный момент в заделке у a2, Н·м", -WA["MB"] / 1000, 6.98, 0.006)
chk("местный момент под камнем, Н·м", WA["Mmid"] / 1000, 4.87, 0.01)
chk("сечение А–А при φу, МПа", WA["s"], 93.6, 0.05); chk("сечение А–А при φк, МПа", WB["s"], 74.3, 0.05)
lo, hi = w + 1, w + 80
for _ in range(200):
    m = (lo + hi) / 2
    if max(wall(A, m)["s"], wall(B, m)["s"]) > 98: lo = m
    else: hi = m
chk("H_треб (Ст3сп, w = 16)", lo, 35.5, 0.05)

# ------------------------------------------------------------------ перегрузка
Fov = 120 * L * cos(rad(phk)) ** 3 / e
chk("перегрузка F, Н", Fov, 3758.0, 0.5); chk("k = F/Fук", Fov / Fk, 1.90, 0.005)

# ------------------------------------------------------------------ сверка с документом
doc = Document(DOCX)
FULL = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
i0, i1 = FULL.index("12. Прочностной расчёт звеньев"), FULL.index("13. Сводка принятых параметров")
sec = FULL[i0:i1]
need = ["128,1", "127,0", "17,4", "91,5", "93,6", "26,3", "67,3", "8,97", "35,5 → 36", "17,4 → 18", "3758", "1,90", "0,718",
        "0,0472", "199,1", "40,75", "2515", "2432", "1974", "λ = 40", "h = 22, H = 40", "0,659", "0,807", "16,8", "Ø16 вместо Ø18"]
miss = [s_ for s_ in need if s_ not in sec and s_ not in FULL]
print("в тексте раздела 12 не найдено:", miss or "—")
bad += miss
old = [s_ for s_ in ("сталь 45 (σ_т = 355 МПа) с запасом", "[p] = 30 МПа", "18 × 18 × 10", "0,833", "предыдущ", "прошл", "в отличие от")
       if s_ in FULL]
print("следы прежнего расчёта / упоминания версий:", old or "—")
bad += old
print("ИТОГ:", "расхождений нет" if not bad else f"расхождения: {bad}")
